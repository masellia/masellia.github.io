#!/usr/bin/env python3
"""Build a local annual citation workbook from INSPIRE and Google Scholar."""

from __future__ import annotations

import argparse
import json
import re
import time
import unicodedata
from collections import defaultdict
from datetime import datetime, timezone
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, urljoin, urlparse

import requests
import yaml
from bs4 import BeautifulSoup
from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parent
PROJECT_ROOT = ROOT.parent
CACHE_DIR = ROOT / "cache"
OUTPUT_PATH = ROOT / "citation_tracker.xlsx"
INSPIRE_AUTHOR_ID = "1189872"
INSPIRE_PROFILE = "https://inspirehep.net/authors/1189872?ui-citation-summary=true"
INSPIRE_API = "https://inspirehep.net/api/literature"
SCHOLAR_USER_ID = "fUBWIzgAAAAJ"
SCHOLAR_PROFILE = f"https://scholar.google.com/citations?user={SCHOLAR_USER_ID}&hl=en"
USER_AGENT = (
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36"
)


class ScholarRateLimit(RuntimeError):
    pass


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--source",
        choices=("all", "inspire", "scholar", "workbook"),
        default="all",
        help="Source to refresh; cached data is used for the other source.",
    )
    parser.add_argument("--refresh", action="store_true", help="Ignore cached source responses.")
    parser.add_argument(
        "--scholar-delay",
        type=float,
        default=10.0,
        help="Seconds between Google Scholar detail requests (default: 10).",
    )
    parser.add_argument(
        "--scholar-max-details",
        type=int,
        default=5,
        help="Maximum Scholar article-detail requests per run (default: 5).",
    )
    parser.add_argument(
        "--scholar-cache-days",
        type=int,
        default=30,
        help="Refresh cached Scholar details older than this many days (default: 30).",
    )
    parser.add_argument(
        "--scholar-refresh-profile",
        action="store_true",
        help="Refresh the two Scholar profile pages before processing details.",
    )
    parser.add_argument(
        "--scholar-no-details",
        action="store_true",
        help="Refresh Scholar profile matching without requesting article detail pages.",
    )
    return parser.parse_args()


def cache_path(*parts: str) -> Path:
    path = CACHE_DIR.joinpath(*parts)
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def read_json(path: Path, default: Any = None) -> Any:
    if not path.exists():
        return default
    with path.open(encoding="utf-8") as handle:
        return json.load(handle)


def write_json(path: Path, value: Any) -> None:
    temporary = path.with_suffix(path.suffix + ".tmp")
    with temporary.open("w", encoding="utf-8") as handle:
        json.dump(value, handle, ensure_ascii=False, indent=2, sort_keys=True)
    temporary.replace(path)


def request_json(
    session: requests.Session,
    url: str,
    params: dict[str, Any],
    path: Path,
    refresh: bool,
) -> dict[str, Any]:
    if path.exists() and not refresh:
        return read_json(path)

    for attempt in range(5):
        response = session.get(url, params=params, timeout=90)
        if response.status_code == 200:
            payload = response.json()
            write_json(path, payload)
            return payload
        if response.status_code not in (429, 500, 502, 503, 504):
            response.raise_for_status()
        time.sleep(2 ** attempt)
    response.raise_for_status()
    raise RuntimeError(f"Unable to retrieve {response.url}")


def request_scholar_html(
    session: requests.Session, url: str, path: Path, refresh: bool
) -> str:
    if path.exists() and not refresh:
        return path.read_text(encoding="utf-8")
    response = session.get(url, timeout=60)
    if response.status_code == 429:
        raise ScholarRateLimit("Google Scholar returned HTTP 429")
    response.raise_for_status()
    if "unusual traffic" in response.text.lower():
        raise ScholarRateLimit("Google Scholar returned an unusual-traffic page")
    path.write_text(response.text, encoding="utf-8")
    return response.text


def normalize_title(value: str) -> str:
    value = unicodedata.normalize("NFKD", value or "")
    value = "".join(character for character in value if not unicodedata.combining(character))
    value = value.lower().replace("&", " and ")
    return re.sub(r"[^a-z0-9]+", " ", value).strip()


def title_variants(value: str) -> list[str]:
    variants = [value]
    without_prefix = re.sub(r"^arxiv:\s*", "", value, flags=re.I)
    variants.append(without_prefix)
    variants.append(re.sub(r"\s*\(\d{4}\)\s*$", "", without_prefix))
    variants.append(
        re.sub(r":\s*[A-Z][A-Za-z. -]+\bet\s+al\.?\s*$", "", without_prefix)
    )
    return list(dict.fromkeys(normalize_title(variant) for variant in variants if variant))


def normalize_arxiv(value: str | None) -> str | None:
    if not value:
        return None
    value = re.sub(r"^arxiv:\s*", "", value.strip(), flags=re.I)
    return re.sub(r"v\d+$", "", value)


def year_from_arxiv(arxiv_id: str | None) -> int | None:
    value = normalize_arxiv(arxiv_id)
    if not value:
        return None
    modern = re.match(r"^(\d{2})(\d{2})\.\d+$", value)
    if modern:
        return 2000 + int(modern.group(1))
    legacy = re.match(r"^[a-z-]+/(\d{2})(\d{2})\d+$", value, flags=re.I)
    if legacy:
        short_year = int(legacy.group(1))
        return 1900 + short_year if short_year >= 91 else 2000 + short_year
    return None


def month_from_arxiv(arxiv_id: str | None) -> int:
    value = normalize_arxiv(arxiv_id)
    if not value:
        return 0
    modern = re.match(r"^\d{2}(\d{2})\.\d+$", value)
    if modern:
        return int(modern.group(1))
    legacy = re.match(r"^[a-z-]+/\d{2}(\d{2})\d+$", value, flags=re.I)
    return int(legacy.group(1)) if legacy else 0


def metadata_title(metadata: dict[str, Any]) -> str:
    titles = metadata.get("titles") or []
    return titles[0].get("title", "") if titles else ""


def metadata_arxiv(metadata: dict[str, Any]) -> str | None:
    values = metadata.get("arxiv_eprints") or []
    return normalize_arxiv(values[0].get("value")) if values else None


def metadata_arxiv_year(metadata: dict[str, Any]) -> int | None:
    arxiv_id = metadata_arxiv(metadata)
    if not arxiv_id:
        return None
    preprint_date = metadata.get("preprint_date")
    if preprint_date and re.match(r"^\d{4}", str(preprint_date)):
        return int(str(preprint_date)[:4])
    return year_from_arxiv(arxiv_id)


def load_website_targets() -> list[dict[str, Any]]:
    path = PROJECT_ROOT / "_data" / "publications.yml"
    data = yaml.safe_load(path.read_text(encoding="utf-8")) or []
    targets = []
    for group in data:
        for paper in group.get("papers") or []:
            arxiv_id = normalize_arxiv(paper.get("arxiv"))
            if arxiv_id:
                targets.append(
                    {
                        "title": paper.get("title", ""),
                        "arxiv": arxiv_id,
                        "target_year": year_from_arxiv(arxiv_id) or group.get("year"),
                    }
                )
    return targets


def collect_inspire(refresh: bool) -> dict[str, Any]:
    session = requests.Session()
    session.headers.update({"User-Agent": USER_AGENT, "Accept": "application/json"})
    fields = ",".join(
        (
            "control_number",
            "titles",
            "arxiv_eprints",
            "preprint_date",
            "earliest_date",
            "citation_count",
            "citation_count_without_self_citations",
            "document_type",
            "dois",
        )
    )
    author_payload = request_json(
        session,
        INSPIRE_API,
        {
            "q": f"authors.recid:{INSPIRE_AUTHOR_ID}",
            "sort": "mostrecent",
            "size": 1000,
            "page": 1,
            "fields": fields,
        },
        cache_path("inspire", "author_literature.json"),
        refresh,
    )
    authored_hits = author_payload.get("hits", {}).get("hits", [])
    papers: list[dict[str, Any]] = []
    raw_citations: list[dict[str, Any]] = []

    citation_fields = ",".join(
        (
            "control_number",
            "titles",
            "arxiv_eprints",
            "preprint_date",
            "earliest_date",
            "dois",
        )
    )
    for number, hit in enumerate(authored_hits, start=1):
        metadata = hit.get("metadata", {})
        source_id = str(metadata.get("control_number") or hit.get("id"))
        title = metadata_title(metadata)
        arxiv_id = metadata_arxiv(metadata)
        target_year = metadata_arxiv_year(metadata)
        source_total = int(metadata.get("citation_count") or 0)
        annual: dict[int, int] = defaultdict(int)
        unresolved = 0
        seen: set[str] = set()

        page = 1
        while True:
            payload = request_json(
                session,
                INSPIRE_API,
                {
                    "q": f"refersto:recid:{source_id}",
                    "sort": "mostrecent",
                    "size": 1000,
                    "page": page,
                    "fields": citation_fields,
                },
                cache_path("inspire", "citations", f"{source_id}_page_{page}.json"),
                refresh,
            )
            citation_hits = payload.get("hits", {}).get("hits", [])
            for citing_hit in citation_hits:
                citing_metadata = citing_hit.get("metadata", {})
                citing_id = str(citing_metadata.get("control_number") or citing_hit.get("id"))
                if citing_id in seen:
                    continue
                seen.add(citing_id)
                citing_year = metadata_arxiv_year(citing_metadata)
                assigned_year = None
                posteriori = False
                if target_year is not None and citing_year is not None:
                    assigned_year = max(target_year, citing_year)
                    posteriori = citing_year < target_year
                    annual[assigned_year] += 1
                else:
                    unresolved += 1
                raw_citations.append(
                    {
                        "target_source_id": source_id,
                        "target_title": title,
                        "target_arxiv": arxiv_id,
                        "target_year": target_year,
                        "citing_source_id": citing_id,
                        "citing_title": metadata_title(citing_metadata),
                        "citing_arxiv": metadata_arxiv(citing_metadata),
                        "citing_year": citing_year,
                        "assigned_year": assigned_year,
                        "a_posteriori": posteriori,
                    }
                )
            total = int(payload.get("hits", {}).get("total") or 0)
            if page * 1000 >= total:
                break
            page += 1
        retrieved_total = len(seen)
        papers.append(
            {
                "source_id": source_id,
                "title": title,
                "arxiv": arxiv_id,
                "target_year": target_year,
                "annual": {str(year): count for year, count in sorted(annual.items())},
                "unresolved": unresolved,
                "source_total": source_total,
                "retrieved_total": retrieved_total,
                "total_matches": source_total == retrieved_total,
            }
        )
        if number == 1 or number % 10 == 0 or number == len(authored_hits):
            print(f"INSPIRE {number}/{len(authored_hits)}: {title[:65]} ({retrieved_total})", flush=True)

    result = {
        "collected_at": datetime.now(timezone.utc).isoformat(),
        "profile_url": INSPIRE_PROFILE,
        "papers": papers,
        "raw_citations": raw_citations,
    }
    write_json(cache_path("processed", "inspire.json"), result)
    return result


def scholar_publication_id(href: str) -> str:
    query = parse_qs(urlparse(href).query)
    return (query.get("citation_for_view") or [""])[0]


def parse_scholar_profile(html: str) -> list[dict[str, Any]]:
    soup = BeautifulSoup(html, "html.parser")
    publications = []
    for row in soup.select("tr.gsc_a_tr"):
        title_link = row.select_one("a.gsc_a_at")
        if not title_link:
            continue
        href = urljoin("https://scholar.google.com", title_link.get("href", ""))
        publication_id = scholar_publication_id(href)
        count_link = row.select_one("td.gsc_a_c a")
        count_text = count_link.get_text(strip=True) if count_link else ""
        publications.append(
            {
                "source_id": publication_id,
                "title": title_link.get_text(" ", strip=True),
                "detail_url": href,
                "source_total": int(count_text) if count_text.isdigit() else 0,
                "display_year": (
                    row.select_one("span.gsc_a_hyear").get_text(strip=True)
                    if row.select_one("span.gsc_a_hyear")
                    else None
                ),
            }
        )
    return publications


def parse_scholar_profile_total(html: str) -> int | None:
    cells = [cell.get_text(strip=True) for cell in BeautifulSoup(html, "html.parser").select("#gsc_rsb_st td")]
    return int(cells[1]) if len(cells) > 1 and cells[1].isdigit() else None


def parse_scholar_profile_annual(html: str) -> dict[str, int]:
    soup = BeautifulSoup(html, "html.parser")
    years = [item.get_text(strip=True) for item in soup.select(".gsc_g_t")]
    counts = [item.get_text(strip=True) for item in soup.select(".gsc_g_a")]
    if len(years) != len(counts):
        return {}
    return {
        year: int(count)
        for year, count in zip(years, counts)
        if year.isdigit() and count.isdigit()
    }


def parse_scholar_annual(html: str) -> dict[int, int]:
    soup = BeautifulSoup(html, "html.parser")
    annual: dict[int, int] = {}
    for bar in soup.select("a.gsc_oci_g_a"):
        label = bar.select_one("span.gsc_oci_g_al")
        query = parse_qs(urlparse(bar.get("href", "")).query)
        year_text = (query.get("as_ylo") or [""])[0]
        count_text = label.get_text(strip=True) if label else ""
        if year_text.isdigit() and count_text.isdigit():
            annual[int(year_text)] = int(count_text)
    return annual


def build_target_index(inspire: dict[str, Any] | None) -> tuple[dict[str, dict[str, Any]], list[dict[str, Any]]]:
    targets = load_website_targets()
    if inspire:
        for paper in inspire.get("papers", []):
            if paper.get("arxiv"):
                targets.append(
                    {
                        "title": paper["title"],
                        "arxiv": paper["arxiv"],
                        "target_year": paper["target_year"],
                    }
                )
    exact: dict[str, dict[str, Any]] = {}
    unique: dict[str, dict[str, Any]] = {}
    for target in targets:
        key = normalize_title(target["title"])
        exact[key] = target
        unique[target["arxiv"]] = target
    return exact, list(unique.values())


def match_target(
    title: str,
    exact: dict[str, dict[str, Any]],
    targets: list[dict[str, Any]],
) -> tuple[dict[str, Any] | None, float]:
    keys = title_variants(title)
    for key in keys:
        if key in exact:
            return exact[key], 1.0
    candidates = sorted(
        (
            (
                max(
                    SequenceMatcher(None, key, normalize_title(target["title"])).ratio()
                    for key in keys
                ),
                target,
            )
            for target in targets
        ),
        key=lambda item: item[0],
        reverse=True,
    )
    if not candidates or candidates[0][0] < 0.92:
        return None, candidates[0][0] if candidates else 0.0
    if len(candidates) > 1 and candidates[0][0] - candidates[1][0] < 0.025:
        return None, candidates[0][0]
    return candidates[0][1], candidates[0][0]


def collect_scholar(
    refresh_profile: bool,
    delay: float,
    inspire: dict[str, Any] | None,
    fetch_details: bool = True,
    max_detail_requests: int = 5,
    detail_cache_days: int = 30,
) -> dict[str, Any]:
    session = requests.Session()
    session.headers.update({"User-Agent": USER_AGENT, "Accept-Language": "en-US,en;q=0.9"})
    publications: list[dict[str, Any]] = []
    profile_total = None
    profile_annual: dict[str, int] = {}
    for start in (0, 100):
        url = f"{SCHOLAR_PROFILE}&cstart={start}&pagesize=100"
        html = request_scholar_html(
            session,
            url,
            cache_path("scholar", f"profile_{start}.html"),
            refresh_profile,
        )
        if start == 0:
            profile_total = parse_scholar_profile_total(html)
            profile_annual = parse_scholar_profile_annual(html)
        publications.extend(parse_scholar_profile(html))

    seen: set[str] = set()
    publications = [
        publication
        for publication in publications
        if publication["source_id"] and not (publication["source_id"] in seen or seen.add(publication["source_id"]))
    ]
    exact, targets = build_target_index(inspire)
    result_papers: list[dict[str, Any]] = []
    annual_raw: list[dict[str, Any]] = []
    rate_limited = False
    detail_requests_made = 0

    def detail_path(publication: dict[str, Any]) -> Path:
        return cache_path(
            "scholar",
            "details",
            re.sub(r"[^A-Za-z0-9_.-]+", "_", publication["source_id"]) + ".html",
        )

    now = time.time()
    max_age_seconds = max(0, detail_cache_days) * 86400
    missing = sorted(
        (publication for publication in publications if not detail_path(publication).exists()),
        key=lambda publication: -publication["source_total"],
    )
    stale = sorted(
        (
            publication
            for publication in publications
            if detail_path(publication).exists()
            and now - detail_path(publication).stat().st_mtime > max_age_seconds
        ),
        key=lambda publication: detail_path(publication).stat().st_mtime,
    )
    request_candidates = {
        publication["source_id"]
        for publication in (missing + stale)[: max(0, max_detail_requests)]
    }

    for number, publication in enumerate(publications, start=1):
        target, match_score = match_target(publication["title"], exact, targets)
        original_annual: dict[int, int] = {}
        detail_status = "not requested"
        detail_cache = detail_path(publication)
        cache_is_stale = (
            detail_cache.exists() and now - detail_cache.stat().st_mtime > max_age_seconds
        )
        should_request = (
            fetch_details
            and not rate_limited
            and publication["source_id"] in request_candidates
        )
        if should_request:
            detail_requests_made += 1
            try:
                html = request_scholar_html(
                    session, publication["detail_url"], detail_cache, refresh=True
                )
                original_annual = parse_scholar_annual(html)
                detail_status = "fetched"
                time.sleep(max(0.0, delay))
            except ScholarRateLimit:
                rate_limited = True
                detail_status = "rate limited"
        if detail_cache.exists() and detail_status != "fetched":
            original_annual = parse_scholar_annual(detail_cache.read_text(encoding="utf-8"))
            if detail_status == "rate limited":
                detail_status = "cached stale; refresh rate limited"
            elif cache_is_stale:
                detail_status = "cached stale"
            else:
                detail_status = "cached"
        elif fetch_details and detail_status == "not requested":
            detail_status = "deferred after rate limit" if rate_limited else "deferred by batch limit"

        target_year = target.get("target_year") if target else None
        arxiv_id = target.get("arxiv") if target else None
        adjusted: dict[int, int] = defaultdict(int)
        for source_year, count in sorted(original_annual.items()):
            assigned_year = max(target_year, source_year) if target_year else None
            if assigned_year is not None:
                adjusted[assigned_year] += count
            annual_raw.append(
                {
                    "target_source_id": publication["source_id"],
                    "target_title": publication["title"],
                    "target_arxiv": arxiv_id,
                    "target_year": target_year,
                    "source_year": source_year,
                    "source_count": count,
                    "assigned_year": assigned_year,
                    "a_posteriori_count": count if target_year and source_year < target_year else 0,
                }
            )
        graphed_total = sum(original_annual.values())
        unresolved = max(0, publication["source_total"] - sum(adjusted.values()))
        result_papers.append(
            {
                **publication,
                "arxiv": arxiv_id,
                "target_year": target_year,
                "match_score": round(match_score, 4),
                "matched_title": target.get("title") if target else None,
                "annual": {str(year): count for year, count in sorted(adjusted.items())},
                "unresolved": unresolved,
                "graphed_total": graphed_total,
                "detail_status": detail_status,
                "total_matches": graphed_total == publication["source_total"],
            }
        )
        if number == 1 or number % 10 == 0 or number == len(publications):
            print(
                f"Scholar {number}/{len(publications)}: {publication['title'][:60]} "
                f"({detail_status})",
                flush=True,
            )

    result = {
        "collected_at": datetime.now(timezone.utc).isoformat(),
        "profile_url": SCHOLAR_PROFILE,
        "profile_total": profile_total,
        "profile_annual": profile_annual,
        "details_requested": fetch_details,
        "detail_requests_made": detail_requests_made,
        "detail_request_limit": max_detail_requests,
        "rate_limited": rate_limited,
        "papers": result_papers,
        "annual_raw": annual_raw,
    }
    write_json(cache_path("processed", "scholar.json"), result)
    return result


def all_years(*datasets: dict[str, Any] | None) -> list[int]:
    years: set[int] = set()
    for dataset in datasets:
        if not dataset:
            continue
        for paper in dataset.get("papers", []):
            years.update(int(year) for year in paper.get("annual", {}))
            if paper.get("target_year"):
                years.add(int(paper["target_year"]))
    if not years:
        return [datetime.now().year]
    return list(range(min(years), max(max(years), datetime.now().year) + 1))


GOLD = "C99A22"
CHARCOAL = "252525"
LIGHT = "F2F1ED"
WHITE = "FFFFFF"
RULE = Side(style="thin", color="D6D3CB")


def style_header(ws, row: int, columns: int) -> None:
    for cell in ws[row][:columns]:
        cell.fill = PatternFill("solid", fgColor=CHARCOAL)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def consolidate_resolved_papers(papers: list[dict[str, Any]]) -> list[dict[str, Any]]:
    consolidated: dict[str, dict[str, Any]] = {}
    for paper in papers:
        if not paper.get("arxiv") or not paper.get("target_year"):
            continue
        key = paper["arxiv"]
        if key not in consolidated:
            consolidated[key] = {
                **paper,
                "title": paper.get("matched_title") or paper["title"],
                "annual": dict(paper.get("annual", {})),
            }
            continue
        existing = consolidated[key]
        for year, count in paper.get("annual", {}).items():
            existing["annual"][year] = int(existing["annual"].get(year, 0)) + int(count)
        existing["unresolved"] = int(existing.get("unresolved") or 0) + int(paper.get("unresolved") or 0)
        existing["source_total"] = int(existing.get("source_total") or 0) + int(paper.get("source_total") or 0)
        existing["source_id"] = f"{existing['source_id']}; {paper['source_id']}"
    return list(consolidated.values())


def source_year_totals(
    dataset: dict[str, Any], resolved: list[dict[str, Any]], years: list[int]
) -> tuple[list[int], int]:
    derived = {
        year: sum(int(paper.get("annual", {}).get(str(year), 0)) for paper in resolved)
        for year in years
    }
    profile_annual = {
        int(year): int(count) for year, count in dataset.get("profile_annual", {}).items()
    }
    if profile_annual:
        first_profile_year = min(profile_annual)
        totals = {
            year: profile_annual.get(year, derived[year] if year < first_profile_year else 0)
            for year in years
        }
        for record in dataset.get("annual_raw", []):
            count = int(record.get("a_posteriori_count") or 0)
            source_year = record.get("source_year")
            assigned_year = record.get("assigned_year")
            if count and source_year in profile_annual and assigned_year in totals:
                totals[source_year] -= count
                totals[assigned_year] += count
    else:
        totals = derived
    grand_total = sum(int(paper.get("source_total") or 0) for paper in dataset.get("papers", []))
    unresolved_total = grand_total - sum(totals.values())
    return [totals[year] for year in years], unresolved_total


def add_citation_charts(
    ws, years: list[int], header_row: int, annual_start: int, annual_end: int, anchor_column: int
) -> None:
    categories = Reference(ws, min_col=annual_start, max_col=annual_end, min_row=header_row)
    charts = (
        (1, "Citations per year", "Citations"),
        (3, "Cumulative citations", "Total citations"),
    )
    for index, (data_row, title, y_title) in enumerate(charts):
        chart = LineChart()
        chart.title = title
        chart.style = 13
        chart.y_axis.title = y_title
        chart.x_axis.title = "Year"
        chart.height = 7
        chart.width = 13
        chart.legend = None
        chart.visible_cells_only = False
        data = Reference(ws, min_col=annual_start, max_col=annual_end, min_row=data_row)
        chart.add_data(data, from_rows=True, titles_from_data=False)
        chart.set_categories(categories)
        series = chart.series[0]
        series.graphicalProperties.line.solidFill = GOLD
        series.graphicalProperties.line.width = 24000
        series.marker.symbol = "circle"
        series.marker.size = 5
        anchor = f"{get_column_letter(anchor_column)}{1 + index * 17}"
        ws.add_chart(chart, anchor)


def add_source_sheet(wb: Workbook, name: str, dataset: dict[str, Any], years: list[int]) -> None:
    ws = wb.create_sheet(name)
    headers = ["arXiv year", "arXiv ID", "Title", *years, "Unresolved year", "Total"]
    resolved = consolidate_resolved_papers(dataset.get("papers", []))
    resolved.sort(
        key=lambda paper: (
            -int(paper["target_year"]),
            -month_from_arxiv(paper["arxiv"]),
            paper["title"].lower(),
        )
    )
    year_totals, unresolved_total = source_year_totals(dataset, resolved, years)
    grand_total = sum(int(paper.get("source_total") or 0) for paper in dataset.get("papers", []))
    ws.append([None, None, "OVERALL TOTAL", *year_totals, unresolved_total, grand_total])
    current_year = datetime.now().year
    current_month = datetime.now().month
    monthly_rates = [
        total / (current_month if year == current_year else 12)
        for year, total in zip(years, year_totals)
    ]
    ws.append([None, None, "CITATIONS / MONTH", *monthly_rates, None, None])
    cumulative = []
    running_total = 0
    for total in year_totals:
        running_total += total
        cumulative.append(running_total)
    ws.append([None, None, "CUMULATIVE CITATIONS", *cumulative, None, None])
    ws.row_dimensions[3].hidden = True
    ws.append(headers)
    header_row = ws.max_row
    style_header(ws, header_row, len(headers))

    for paper in resolved:
        annual = paper.get("annual", {})
        row = [paper["target_year"], paper["arxiv"], paper["title"]]
        row.extend(int(annual.get(str(year), 0)) for year in years)
        row.extend((int(paper.get("unresolved") or 0), int(paper.get("source_total") or 0)))
        ws.append(row)

    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=LIGHT)
        cell.font = Font(bold=True, color=CHARCOAL)
        cell.border = Border(bottom=Side(style="medium", color=GOLD))
    for cell in ws[2]:
        cell.font = Font(italic=True, color=CHARCOAL)
    for cell in ws[2][3 : 3 + len(years)]:
        cell.number_format = "0.00"
        cell.alignment = Alignment(horizontal="center")

    ws.freeze_panes = "D5"
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{ws.max_row}"
    ws.column_dimensions["C"].width = 58
    annual_start = 4
    annual_end = annual_start + len(years) - 1
    for column in range(annual_start, annual_end + 1):
        ws.column_dimensions[get_column_letter(column)].width = 6
    for column in (1, 2, annual_end + 1, annual_end + 2):
        longest = max(
            len(str(ws.cell(row, column).value))
            for row in range(1, ws.max_row + 1)
            if ws.cell(row, column).value is not None
        )
        ws.column_dimensions[get_column_letter(column)].width = longest + 1
    for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row):
        for cell in row:
            cell.border = Border(bottom=RULE)
        row[2].alignment = Alignment(wrap_text=True, vertical="top")
    add_citation_charts(
        ws,
        years,
        header_row,
        annual_start,
        annual_end,
        len(headers) + 2,
    )
    ws.sheet_view.showGridLines = False


def add_metadata_sheet(
    wb: Workbook, inspire: dict[str, Any] | None, scholar: dict[str, Any] | None
) -> None:
    ws = wb.create_sheet("Metadata", 0)
    rows = [
        ("Generated", datetime.now(timezone.utc).isoformat()),
        ("INSPIRE profile", INSPIRE_PROFILE),
        ("Google Scholar profile", SCHOLAR_PROFILE),
        ("Citation policy", "All citations are included, including self-citations."),
        (
            "INSPIRE total basis",
            "Sum of all author-profile literature records; the UI citeable summary may exclude non-citeable books or proceedings.",
        ),
        ("Paper year", "First arXiv submission year; journal publication year is not used."),
        (
            "A posteriori rule",
            "Assigned year = max(cited-paper arXiv year, citing-paper/source citation year).",
        ),
        (
            "Unresolved year",
            "Includes citations lacking an arXiv-based year and all citations to papers without a matched arXiv ID.",
        ),
    ]
    if inspire:
        rows.append(("INSPIRE collected", inspire.get("collected_at")))
        rows.append(("INSPIRE papers", len(inspire.get("papers", []))))
    if scholar:
        rows.append(("Scholar collected", scholar.get("collected_at")))
        rows.append(("Scholar papers", len(scholar.get("papers", []))))
        rows.append(("Scholar profile total", scholar.get("profile_total")))
        rows.append(("Scholar profile annual range", ", ".join(scholar.get("profile_annual", {}).keys())))
        rows.append(("Scholar details requested", scholar.get("details_requested")))
        rows.append(("Scholar detail requests made", scholar.get("detail_requests_made")))
        rows.append(("Scholar detail request limit", scholar.get("detail_request_limit")))
        rows.append(("Scholar rate limited", scholar.get("rate_limited")))
    for key, value in rows:
        ws.append([key, value])
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 110
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=CHARCOAL)
        cell.font = Font(color=WHITE, bold=True)
    for row in ws.iter_rows():
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    ws.sheet_view.showGridLines = False


def build_workbook(inspire: dict[str, Any] | None, scholar: dict[str, Any] | None) -> None:
    if not inspire and not scholar:
        raise RuntimeError("No cached source data is available. Collect a source first.")
    years = all_years(inspire, scholar)
    wb = Workbook()
    wb.remove(wb.active)
    add_metadata_sheet(wb, inspire, scholar)
    if inspire:
        add_source_sheet(wb, "INSPIRE", inspire, years)
    if scholar:
        add_source_sheet(wb, "Google Scholar", scholar, years)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.save(OUTPUT_PATH)
    print(f"Workbook written to {OUTPUT_PATH}")


def validate_workbook(inspire: dict[str, Any] | None, scholar: dict[str, Any] | None) -> None:
    workbook = load_workbook(OUTPUT_PATH, data_only=False, read_only=True)
    expected = {"Metadata"}
    if inspire:
        expected.add("INSPIRE")
    if scholar:
        expected.add("Google Scholar")
    missing = expected.difference(workbook.sheetnames)
    if missing:
        raise RuntimeError(f"Workbook is missing sheets: {', '.join(sorted(missing))}")
    for name, dataset in (("INSPIRE", inspire), ("Google Scholar", scholar)):
        if not dataset:
            continue
        expected_total = sum(int(paper.get("source_total") or 0) for paper in dataset.get("papers", []))
        if dataset.get("profile_total") is not None and dataset["profile_total"] != expected_total:
            raise RuntimeError(
                f"{name} paper sum {expected_total} != profile total {dataset['profile_total']}"
            )
        sheet = workbook[name]
        actual_total = sheet.cell(1, sheet.max_column).value
        if actual_total != expected_total:
            raise RuntimeError(f"{name} workbook total {actual_total} != source total {expected_total}")
    print("Workbook structure and source grand totals validated")


def main() -> None:
    args = parse_args()
    inspire = read_json(cache_path("processed", "inspire.json"))
    scholar = read_json(cache_path("processed", "scholar.json"))
    if args.source in ("all", "inspire"):
        inspire = collect_inspire(args.refresh)
    if args.source in ("all", "scholar"):
        try:
            scholar = collect_scholar(
                args.refresh or args.scholar_refresh_profile,
                args.scholar_delay,
                inspire,
                fetch_details=not args.scholar_no_details,
                max_detail_requests=args.scholar_max_details,
                detail_cache_days=args.scholar_cache_days,
            )
        except ScholarRateLimit as error:
            print(f"Scholar collection stopped before profile retrieval: {error}")
    build_workbook(inspire, scholar)
    validate_workbook(inspire, scholar)


if __name__ == "__main__":
    main()
